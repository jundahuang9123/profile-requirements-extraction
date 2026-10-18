"""Offline checks for pinned base data and its RDF/DOCX parsers."""
import importlib.util
import json
import hashlib
from pathlib import Path

import pytest
from docx import Document
from rdflib import Graph, URIRef

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('base_registry_builder', ROOT / 'scripts/build_base_profile.py')
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
REGISTRY = ROOT / 'requirement-reuse-service/config/base_profiles'
PREFIXES = '''
@prefix : <https://example.test/> .
@prefix sh: <http://www.w3.org/ns/shacl#> .
@prefix dcat: <http://www.w3.org/ns/dcat#> .
@prefix dct: <http://purl.org/dc/terms/> .
@prefix foaf: <http://xmlns.com/foaf/0.1/> .
@prefix skos: <http://www.w3.org/2004/02/skos/core#> .
@prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
'''


@pytest.mark.parametrize('identifier', builder.SUPPORTED_PROFILES)
def test_every_registry_is_bound_to_official_pinned_input_bytes(identifier):
    data = json.loads((REGISTRY / f'{identifier}.json').read_text())
    assert data['id'] == identifier
    for path, digest in data['source']['files'].items():
        assert digest == builder.PINNED_HASHES[path]
    if identifier == 'dcat-ap-1.1.0':
        assert data['source']['url'] == 'https://datos.gob.es/sites/default/files/doc/file/dcat-ap_version_1.1.pdf'
    elif identifier.startswith('dcat-ap-'):
        assert data['source']['commit'] == builder.PINNED_COMMIT
    else:
        assert data['source']['url'] == builder.W3C_SOURCES[data['version']]


@pytest.mark.parametrize('version', ['3.0.0', '3.0.1'])
def test_ap3_registry_combines_separate_constraints_and_keeps_vocabularies(version):
    data = json.loads((REGISTRY / f'dcat-ap-{version}.json').read_text())
    classes = data['classes']
    assert classes['dcat:Dataset']['dct:title']['obligation'] == 'mandatory'
    assert classes['dcat:Dataset']['dct:publisher']['obligation'] == 'recommended'
    assert classes['dcat:Dataset']['dct:publisher']['max'] == 1
    assert classes['dcat:Dataset']['dct:publisher']['range'] == 'foaf:Agent'
    assert classes['dcat:DatasetSeries']['dct:title']['obligation'] == 'mandatory'
    assert classes['dcat:Dataset']['dct:language']['vocabulary'] == 'http://publications.europa.eu/resource/authority/language'
    assert classes['dcat:Dataset']['dcat:theme']['vocabulary'] == 'http://publications.europa.eu/resource/authority/data-theme'
    assert classes['dcat:Dataset']['dct:publisher']['vocabulary'] is None
    assert classes['dcat:Dataset']['dct:publisher']['vocabulary_guidance']['qualification'] == 'conditional'


@pytest.mark.parametrize('version', ['2', '3'])
def test_w3c_vocabulary_inherits_properties_without_fabricating_obligations(version):
    data = json.loads((REGISTRY / f'dcat-{version}.0.0.json').read_text())
    for properties in data['classes'].values():
        assert all(entry['obligation'] is None and entry['max'] is None for entry in properties.values())
    assert 'dct:title' in data['classes']['dcat:Dataset']
    assert ('dcat:DatasetSeries' in data['classes']) == (version == '3')
    if version == '3':
        assert 'dct:title' in data['classes']['dcat:DatasetSeries']


def test_ap11_uses_pinned_final_pdf_for_property_tables_and_vocabularies():
    data = builder.build_profile('dcat-ap-1.1.0', REGISTRY / 'source-inputs')
    classes = data['classes']
    assert data['source']['files']['releases/1.1/dcat-ap_1.1.pdf'] == \
        '7d1ab3eeef3fd832e3b289b664e9dedf61197c348a933d3ea0a9b7d685b9124d'
    assert sum(map(len, classes.values())) == 78
    assert classes['dcat:Catalog']['dct:title']['obligation'] == 'mandatory'
    assert classes['dcat:Catalog']['dct:title']['max'] is None
    assert 'dct:titleCatalogue' not in classes['dcat:Catalog']
    assert classes['dcat:Dataset']['dct:publisher']['vocabulary'] is None
    assert classes['dcat:Dataset']['dct:publisher']['vocabulary_guidance']['qualification'] == 'conditional'
    assert classes['dcat:Distribution']['dcat:mediaType']['vocabulary'] == \
        'http://www.iana.org/assignments/media-types/media-types.xhtml'
    assert classes['dcat:Dataset']['dcat:theme']['vocabulary'] == \
        'http://publications.europa.eu/resource/authority/data-theme'
    assert classes['dcat:Dataset']['dct:title']['usage_note_source']['page'] == 13
    assert all(not any(note.strip().startswith(('0..', '1..', '2..')) for note in entry.get('usage_note', '').splitlines())
               for props in classes.values() for entry in props.values())


def test_ap11_needs_catalogue_preserves_pdf_usage_note_with_locator():
    needs = json.loads((REGISTRY / 'dcat-ap-1.1.0.needs.json').read_text())
    entry = next(row for row in needs['entries']
                 if row.get('class') == 'dcat:Dataset' and row.get('property') == 'dct:title')
    assert entry['usage_notes']
    assert any('repeated for parallel language versions' in note for note in entry['usage_notes'])
    assert entry['definition_sources'][0]['file'] == 'releases/1.1/dcat-ap_1.1.pdf'
    assert entry['definition_sources'][0]['page'] == 13


def pinned_fixture(tmp_path, monkeypatch, relative, text):
    path = tmp_path / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text)
    monkeypatch.setitem(builder.PINNED_HASHES, relative, hashlib.sha256(path.read_bytes()).hexdigest())
    return path


def test_ap3_html_parser_and_multiple_shapes_are_not_last_shape_wins(tmp_path, monkeypatch):
    classes = ['Catalog', 'CatalogRecord', 'Dataset', 'DatasetSeries', 'Distribution', 'DataService']
    summary, details = [], []
    for cls in classes:
        summary.append(f'<tr><td>{cls}</td><td>dcat:{cls}</td><td>mandatory</td><td><a href="#{cls}.title">title</a></td><td>dct:title</td></tr>')
        details.append(f'<tr id="{cls}.title"><td>title</td><td>rdfs:Literal</td><td>1..*</td></tr>')
    summary.append('<tr><td>Dataset</td><td>dcat:Dataset</td><td>recommended</td><td><a href="#Dataset.publisher">publisher</a></td><td>dct:publisher</td></tr>')
    details.append('<tr id="Dataset.publisher"><td>publisher</td><td>foaf:Agent</td><td>0..1</td></tr>')
    document = '<html><body><table>' + ''.join(summary + details) + '</table></body></html>'
    pinned_fixture(tmp_path, monkeypatch, 'releases/3.0.0/index.html', document)
    shapes = PREFIXES + '''
    :Dataset sh:targetClass dcat:Dataset; sh:property
      [sh:path dct:publisher; sh:maxCount 1],
      [sh:path dct:publisher; sh:class foaf:Agent],
      [sh:path dct:publisher; sh:nodeKind sh:BlankNodeOrIRI],
      [sh:path dct:title; sh:minCount 1] .
    '''
    path = pinned_fixture(tmp_path, monkeypatch, 'releases/3.0.0/shacl/dcat-ap-SHACL.ttl', shapes)
    data = builder.build_profile('dcat-ap-3.0.0', tmp_path)
    entry = data['classes']['dcat:Dataset']['dct:publisher']
    assert entry['obligation'] == 'recommended' and entry['max'] == 1 and entry['range'] == 'foaf:Agent'
    assert data['classes']['dcat:Dataset']['dct:title']['obligation'] == 'mandatory'
    pinned_fixture(tmp_path, monkeypatch, 'releases/3.0.0/shacl/dcat-ap-SHACL.ttl', shapes.replace('sh:maxCount 1', 'sh:maxCount 2'))
    with pytest.raises(ValueError, match='maximum disagreement'):
        builder.build_profile('dcat-ap-3.0.0', tmp_path)
    path.write_text('Changed source bytes without updating the pin.')
    with pytest.raises(ValueError, match='checksum mismatch'):
        builder.build_profile('dcat-ap-3.0.0', tmp_path)


@pytest.mark.parametrize('version', ['2.0.1', '2.1.1'])
def test_committed_registry_has_pinned_release_contract(version):
    data = json.loads((REGISTRY / f'dcat-ap-{version}.json').read_text())
    assert data['source']['repo'] == 'SEMICeu/DCAT-AP'
    assert data['source']['commit'] == builder.PINNED_COMMIT
    assert data['registry_policy'] == builder.BUILDER_VERSION
    assert data['source']['files']
    for path, digest in data['source']['files'].items():
        assert digest == builder.PINNED_HASHES[path]
        assert len(digest) == 64
    classes = data['classes']
    assert classes['dcat:Dataset']['dct:title']['obligation'] == 'mandatory'
    assert classes['dcat:Dataset']['dct:description']['obligation'] == 'mandatory'
    assert classes['dcat:Dataset']['dct:publisher']['obligation'] == 'recommended'
    assert classes['dcat:Dataset']['dct:publisher']['max'] == 1
    assert classes['dcat:Dataset']['dct:publisher']['range'] == 'foaf:Agent'
    assert classes['dcat:Distribution']['dcat:accessURL']['obligation'] == 'mandatory'
    assert classes['dcat:DataService']['dcat:endpointURL']['obligation'] == 'mandatory'
    assert classes['dcat:DataService']['dct:title']['obligation'] == 'mandatory'
    assert classes['dcat:CatalogRecord']['dct:modified']['obligation'] == 'mandatory'
    assert classes['dcat:Catalog']['dcat:dataset']['obligation'] == ('mandatory' if version == '2.0.1' else 'optional')
    assert 'dcatap:availability' in classes['dcat:Distribution']
    assert 'dct:license' in classes['dcat:Distribution']
    assert 'dct:licence' not in classes['dcat:Distribution']
    assert 'http://data.europa.eu/r5ravailability' not in classes['dcat:Distribution']
    assert 'endpoint address' in classes['dcat:DataService']['dcat:endpointURL']['labels']
    assert 'category' not in classes['dcat:Dataset']['dcat:theme']['labels']
    assert classes['dcat:Dataset']['dcat:theme']['vocabulary'] == 'http://publications.europa.eu/resource/authority/data-theme'
    assert data['recommended_source'] == ('2.0.1-docx' if version == '2.0.1' else '2.1.1-shacl')
    assert 'fallback_note' not in data


def test_rdf_membership_cardinality_ranges_vocabularies_and_typo_normalization():
    shapes = Graph().parse(data=PREFIXES + '''
    :Dataset sh:targetClass dcat:Dataset; sh:property
      [sh:path dct:title; sh:minCount 1; sh:maxCount 3],
      [sh:path dct:publisher; sh:maxCount 1],
      [sh:path dcat:theme] .
    :Distribution sh:targetClass dcat:Distribution; sh:property
      [sh:path <http://data.europa.eu/r5ravailability>],
      [sh:path dct:licence] .
    ''', format='turtle')
    recommended = Graph().parse(data=PREFIXES + '''
    :Dataset sh:targetClass dcat:Dataset; sh:property [sh:path dct:publisher] .
    ''', format='turtle')
    ranges = Graph().parse(data=PREFIXES + '''
    :Dataset sh:targetClass dcat:Dataset; sh:property
      [sh:path dct:publisher; sh:class foaf:Agent],
      [sh:path dct:title; sh:node :String],
      [sh:path dcat:theme; sh:node :Union] .
    :String sh:nodeKind sh:Literal .
    :Union sh:or ([sh:class skos:Concept] [sh:class skos:ConceptScheme]) .
    ''', format='turtle')
    vocabularies = Graph().parse(data=PREFIXES + '''
    :Dataset sh:targetClass dcat:Dataset; sh:property [sh:path dcat:theme; sh:node :Theme] .
    :Theme sh:property [sh:path skos:inScheme; sh:hasValue <https://example.test/themes>] .
    # An erroneous supplement target must not add a property to the base.
    :Unexpected sh:targetClass dcat:Dataset; sh:property [sh:path dcat:accessURL] .
    ''', format='turtle')
    labels = Graph().parse(data=PREFIXES + 'dct:publisher rdfs:label "Dataset - publisher"@en .', format='turtle')
    classes = builder.registry_classes(shapes, recommended=recommended, ranges=ranges, vocabularies=vocabularies, labels=labels)
    assert classes['dcat:Dataset']['dct:title']['obligation'] == 'mandatory'
    assert classes['dcat:Dataset']['dct:title']['max'] == 3
    assert classes['dcat:Dataset']['dct:title']['range'] == 'rdfs:Literal'
    assert classes['dcat:Dataset']['dct:publisher']['obligation'] == 'recommended'
    assert classes['dcat:Dataset']['dct:publisher']['range'] == 'foaf:Agent'
    assert 'publisher' in classes['dcat:Dataset']['dct:publisher']['labels']
    assert classes['dcat:Dataset']['dcat:theme']['range'] == ['skos:Concept', 'skos:ConceptScheme']
    assert classes['dcat:Dataset']['dcat:theme']['vocabulary'] == 'https://example.test/themes'
    assert 'dcat:accessURL' not in classes['dcat:Dataset']
    assert set(classes['dcat:Distribution']) == {'dcatap:availability', 'dct:license'}


def docx_fixture(path):
    doc = Document()
    classes = ['Catalogue', 'Catalogue Record', 'Data Service', 'Dataset', 'Distribution', 'Agent']
    for cls in classes:
        for obligation, uri in [('Mandatory', 'dct:title'), ('Recommended', 'dct:publisher'), ('Optional', 'dct:license')]:
            doc.add_heading(f'{obligation} properties for {cls}', level=3)
            table = doc.add_table(rows=2, cols=5)
            for cell, text in zip(table.rows[0].cells, ['Property', 'URI', 'Range', 'Usage note', 'Card.']):
                cell.text = text
            for cell, text in zip(table.rows[1].cells, [uri.split(':')[1], uri, 'foaf:Agent' if uri == 'dct:publisher' else 'rdfs:Literal', 'Usage', '0..1']):
                cell.text = text
            if cls == 'Dataset' and obligation == 'Recommended':
                row = table.add_row()
                row.cells[0].text = 'theme'
                row.cells[1].text = 'dcat:theme, subproperty of dct:subject'
                row.cells[2].text = 'skos:Concept'
                note = table.add_row()
                note.cells[0].text = 'Please note that both properties are recommended.'
                note.cells[1].text = 'Please note that the order is as follows: dcat:theme.'
    doc.save(path)


def test_docx_reads_heading_association_without_misreading_notes(tmp_path):
    path = tmp_path / 'classes.docx'
    docx_fixture(path)
    properties = builder.parse_docx_properties(path)
    assert properties['dcat:Dataset']['dct:publisher']['obligation'] == 'recommended'
    assert properties['dcat:Dataset']['dct:license']['obligation'] == 'optional'
    assert properties['dcat:Dataset']['dcat:theme']['obligation'] == 'recommended'
    assert len(properties['dcat:Dataset']) == 4
    # DOCX obligations never override a SHACL mandatory count.
    shapes = Graph().parse(data=PREFIXES + ':Dataset sh:targetClass dcat:Dataset; sh:property [sh:path dct:publisher; sh:minCount 1] .', format='turtle')
    classes = builder.registry_classes(shapes, docx_properties=properties)
    assert classes['dcat:Dataset']['dct:publisher']['obligation'] == 'mandatory'
    assert classes['dcat:Dataset']['dct:publisher']['range'] == 'foaf:Agent'


def test_incomplete_docx_is_rejected_instead_of_assuming_optional(tmp_path):
    doc = Document()
    doc.add_heading('Recommended properties for Dataset', level=3)
    path = tmp_path / 'missing-tables.docx'
    doc.save(path)
    with pytest.raises(ValueError, match='incomplete'):
        builder.parse_docx_properties(path)


def test_unpinned_input_is_rejected_before_parse(tmp_path):
    relative = 'releases/2.0.1/dcat-ap_2.0.1_shacl_shapes.ttl'
    file = tmp_path / relative
    file.parent.mkdir(parents=True)
    file.write_text('modified source')
    used = {}
    with pytest.raises(ValueError, match='checksum mismatch'):
        builder.read_pinned(tmp_path, relative, used)
    assert used == {}


def test_vocabulary_parser_does_not_confuse_nested_agent_type_with_publisher_scheme():
    graph = Graph().parse(data=PREFIXES + '''
    :Publisher sh:node :Corporate, :TypedPublisher .
    :Corporate sh:property [sh:path skos:inScheme; sh:hasValue <https://example.test/corporate-bodies>] .
    :TypedPublisher sh:property [sh:path dct:type; sh:node :Types] .
    :Types sh:property [sh:path skos:inScheme; sh:hasValue <https://example.test/publisher-types>] .
    ''', format='turtle')
    assert builder.vocabulary_values(graph, URIRef('https://example.test/Publisher')) == {'https://example.test/corporate-bodies'}


def test_docx_fallback_requires_explicit_opt_in_and_records_condition(tmp_path, monkeypatch):
    minimal_shapes = PREFIXES + ':Dataset sh:targetClass dcat:Dataset; sh:property [sh:path dct:publisher] .'
    files = {
        'releases/2.0.1/dcat-ap_2.0.1_shacl_shapes.ttl': minimal_shapes,
        'releases/2.0.1/dcat-ap_2.0.1_shacl_mdr-vocabularies.shape.ttl': PREFIXES,
        'releases/2.0.1/dcat-ap_2.0.1.rdf': '<rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#"/>',
        'releases/2.0.1/DCAT_AP_2.0.1.docx': 'the mocked parser rejects this fixture',
        'releases/2.1.1/dcat-ap_2.1.1_shacl_shapes_recommended.ttl': minimal_shapes,
    }
    for relative, content in files.items():
        path = tmp_path / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content)
    # The hash gate has its own test; this test isolates fallback policy.
    def read_fixture(input_dir, relative, used):
        used[relative] = builder.PINNED_HASHES[relative]
        return input_dir / relative
    def broken_parse(_path):
        raise ValueError('missing recommended sections')
    monkeypatch.setattr(builder, 'read_pinned', read_fixture)
    monkeypatch.setattr(builder, 'parse_docx_properties', broken_parse)
    with pytest.raises(ValueError, match='no fallback taken'):
        builder.build_profile('dcat-ap-2.0.1', tmp_path)
    data = builder.build_profile('dcat-ap-2.0.1', tmp_path, allow_recommended_fallback=True)
    assert data['recommended_source'] == '2.1.1'
    assert 'missing recommended sections' in data['fallback_note']
    assert 'releases/2.1.1/dcat-ap_2.1.1_shacl_shapes_recommended.ttl' in data['source']['files']
    assert data['classes']['dcat:Dataset']['dct:publisher']['obligation'] == 'recommended'
